import tempfile
import unittest
from pathlib import Path

from docx import Document

from document_generator import DocumentGenerator
from pipeline_whatsapp_docx import run_pipeline


PROJECT_ROOT = Path(__file__).resolve().parents[1]
EXAMPLE_TEXT = (PROJECT_ROOT / "data" / "mensagem_zap.example.txt").read_text(
    encoding="utf-8"
)


class PipelineEndToEndTests(unittest.TestCase):
    def test_generates_three_complete_documents_for_every_area(self) -> None:
        for area in (
            "Estética",
            "Drogaria",
            "Hospitalar",
            "Manipulação",
            "Controle de qualidade",
        ):
            with self.subTest(area=area), tempfile.TemporaryDirectory() as temporary:
                root = Path(temporary)
                input_path = root / "entrada.txt"
                input_path.write_text(
                    EXAMPLE_TEXT.replace("Área estágio: Drogaria", f"Área estágio: {area}"),
                    encoding="utf-8",
                )

                outputs = run_pipeline(input_path, root / "docx", quantidade=3)

                self.assertEqual(3, len(outputs))
                for output in outputs:
                    self.assertTrue(output.is_file())
                    document = Document(output)
                    self.assertFalse(DocumentGenerator._collect_placeholders(document))
                    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                    self.assertNotIn("Texto temporário", text)


if __name__ == "__main__":
    unittest.main()
